"""Convert the documented Markdown subset to an editable research DOCX.

Usage: python scripts/md-to-docx.py input.md --output output.docx
Dependency: python-docx (the application itself needs no Python packages).
Supports headings, paragraphs, inline bold/code/links, ordered/unordered lists,
pipe tables, local PNG/JPEG figures, code fences and editable equation fences.
HTML comments are omitted. Unsupported blocks fail rather than disappear.
Equations use Unicode mathematical notation, not a LaTeX parser.
"""
from __future__ import annotations

import argparse
from pathlib import Path
import re
import sys
from datetime import datetime, timezone
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')


def run_style(run, *, code=False):
    if code:
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0, 0, 0)


def inline(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('`') and part.endswith('`'):
            run_style(paragraph.add_run(part[1:-1]), code=True)
        elif re.fullmatch(r'\[[^\]]+\]\([^)]+\)', part):
            label, target = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', part).groups()
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
            node = OxmlElement('w:r')
            props = OxmlElement('w:rPr')
            color = OxmlElement('w:color')
            color.set(qn('w:val'), '184B70')
            props.append(color)
            node.append(props)
            value = OxmlElement('w:t')
            value.text = label
            node.append(value)
            link.append(node)
            paragraph._p.append(link)
        else:
            paragraph.add_run(part)


def equation(document, value):
    """Native editable Office Math, with structural subscripts and superscripts."""
    paragraph = document.add_paragraph(style='Equation')
    math_para = OxmlElement('m:oMathPara')
    expression = OxmlElement('m:oMath')

    def atom(text):
        node = OxmlElement('m:r')
        t = OxmlElement('m:t')
        t.text = text
        node.append(t)
        return node

    # Unicode superscript squares become actual Office Math superscripts.
    # Underscore identifiers become one structural subscript, e.g. p_XA,i.
    for token in re.split(r'([A-Za-zΔε]+(?:_[A-Za-z]+(?:,[A-Za-z]+)?)?[²]?)', value):
        if not token:
            continue
        match = re.fullmatch(r'([A-Za-zΔε]+)_([A-Za-z]+(?:,[A-Za-z]+)?)', token)
        if match:
            node = OxmlElement('m:sSub')
            base, sub = OxmlElement('m:e'), OxmlElement('m:sub')
            base.append(atom(match[1]))
            sub.append(atom(match[2]))
            node.extend([base, sub])
            expression.append(node)
        elif token.endswith('²'):
            node = OxmlElement('m:sSup')
            base, sup = OxmlElement('m:e'), OxmlElement('m:sup')
            base.append(atom(token[:-1]))
            sup.append(atom('2'))
            node.extend([base, sup])
            expression.append(node)
        else:
            expression.append(atom(token))
    math_para.append(expression)
    paragraph._p.append(math_para)


def cell_value(cell, value):
    paragraph = cell.paragraphs[0]
    paragraph.style = 'Table Text'
    # Deliberate line breaks keep byte identifiers readable in a narrow column.
    if re.fullmatch(r'[0-9a-f]{64}', value):
        for i in range(0, 64, 32):
            if i:
                paragraph.add_run().add_break()
            run = paragraph.add_run(value[i:i + 32])
            run.font.name = 'Consolas'
            run.font.size = Pt(7.4)
    else:
        inline(paragraph, value)


def make_table(document, rows):
    headers = rows[0]
    width = 6.75
    count = len(headers)
    if count == 3:
        fractions = [.22, .25, .53]
    elif count == 4:
        fractions = [.21, .15, .19, .45]
    elif count == 7 and 'Day 8 leaders' in headers:
        fractions = [.09, .09, .08, .08, .22, .22, .22]
    elif count == 8:
        fractions = [.24, .11, .11, .12, .12, .10, .10, .10]
    elif count == 2:
        fractions = [.42, .58]
    else:
        fractions = [1 / count] * count
    table = document.add_table(rows=1, cols=count)
    table.autofit = False
    for col, fraction in zip(table.columns, fractions):
        col.width = Inches(width * fraction)
    props = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        edge = OxmlElement(f'w:{side}')
        edge.set(qn('w:val'), 'single')
        edge.set(qn('w:sz'), '4')
        edge.set(qn('w:color'), 'D0D0D0')
        borders.append(edge)
    props.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for side, size in [('top', 65), ('bottom', 65), ('left', 80), ('right', 80)]:
        edge = OxmlElement(f'w:{side}')
        edge.set(qn('w:w'), str(size))
        edge.set(qn('w:type'), 'dxa')
        margins.append(edge)
    props.append(margins)
    for row_index, values in enumerate(rows):
        row = table.rows[0] if row_index == 0 else table.add_row()
        row_props = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        row_props.append(no_split)
        if row_index == 0:
            row_props.append(OxmlElement('w:tblHeader'))
        for col, value, fraction in zip(row.cells, values, fractions):
            col.width = Inches(width * fraction)
            cell_value(col, value)
            # Keep the header with at least its first data row. Short tables
            # stay intact; longer evidence tables may break with repeated headers.
            if row_index == 0 or (len(rows) <= 7 and row_index < len(rows) - 1):
                col.paragraphs[0].paragraph_format.keep_with_next = True
            if row_index == 0:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'F0F1F2')
                col._tc.get_or_add_tcPr().append(shading)
                for run in col.paragraphs[0].runs:
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def configure(document):
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.left_margin = section.right_margin = Inches(.76)
    section.top_margin = section.bottom_margin = Inches(.70)
    section.header_distance = section.footer_distance = Inches(.30)
    styles = document.styles
    normal = styles['Normal']
    normal.font.name, normal.font.size = 'Cambria', Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in [('Title', 25), ('Subtitle', 13), ('Heading 1', 15), ('Heading 2', 12)]:
        style = styles[name]
        style.font.name, style.font.size = 'Calibri', Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name.startswith('Heading')
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(13 if name.startswith('Heading') else 0)
        style.paragraph_format.space_after = Pt(6)
        props = style.element.find(qn('w:pPr'))
        if props is not None:
            for border in list(props.findall(qn('w:pBdr'))):
                props.remove(border)
    for name, font, size in [('Table Text', 'Calibri', 9), ('Code Block', 'Consolas', 8.7),
                             ('Equation', 'Cambria Math', 11), ('Figure Caption', 'Calibri', 9)]:
        style = styles.add_style(name, 1)
        style.base_style = normal
        style.font.name, style.font.size = font, Pt(size)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.05
    styles['Equation'].paragraph_format.keep_together = True
    styles['Code Block'].paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run('Morrow  •  ')
    run.font.name, run.font.size = 'Calibri', Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)


def convert(source, output):
    content = source.read_text(encoding='utf-8')
    content = re.sub(r'<!--.*?-->', '', content, flags=re.S)
    lines = content.splitlines()
    document = Document()
    configure(document)
    document.core_properties.title = lines[0].removeprefix('# ')
    document.core_properties.author = 'Prepared for Elmar Quliyev'
    document.core_properties.subject = 'Paired artificial life experiments and targeted thrift intervention'
    document.core_properties.keywords = 'Morrow, simulation, reproducibility, artificial life'
    document.core_properties.created = datetime.now(timezone.utc)
    document.core_properties.modified = document.core_properties.created
    document.core_properties.last_modified_by = 'Codex'
    i = 0
    introductory = 0
    stats = {'tables': 0, 'figures': 0, 'equations': 0, 'paragraphs': 0}
    while i < len(lines):
        text = lines[i].strip()
        if not text:
            i += 1
            continue
        if text.startswith('```'):
            language = text[3:]
            end = i + 1
            while end < len(lines) and lines[end].strip() != '```':
                end += 1
            if end == len(lines):
                raise ValueError(f'Unclosed fence at line {i + 1}')
            for value in lines[i + 1:end]:
                if language == 'equation':
                    equation(document, value)
                    stats['equations'] += 1
                else:
                    paragraph = document.add_paragraph(style='Code Block')
                    paragraph.add_run(value)
            if language != 'equation' and end > i + 1:
                paragraph.paragraph_format.keep_with_next = False
            i = end + 1
            continue
        heading = re.fullmatch(r'(#{1,3}) (.+)', text)
        if heading:
            depth, title = len(heading[1]), heading[2]
            if depth == 1 and introductory:
                raise ValueError('Only one document title is supported')
            paragraph = document.add_paragraph(title, style='Title' if depth == 1 else f'Heading {depth - 1}')
            if depth == 1:
                introductory = 1
            else:
                introductory = 3
            i += 1
            continue
        if text.startswith('#'):
            raise ValueError(f'Unsupported heading depth at line {i + 1}')
        image = re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)', text)
        if image:
            path = (source.parent / image[2]).resolve()
            if not path.is_file():
                raise FileNotFoundError(path)
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            picture = paragraph.add_run().add_picture(str(path), width=Inches(6.65))
            picture._inline.docPr.set('descr', image[1])
            stats['figures'] += 1
            i += 1
            continue
        if text.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = [cell.strip() for cell in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', cell) for cell in row):
                    rows.append(row)
                i += 1
            if any(len(row) != len(rows[0]) for row in rows):
                raise ValueError(f'Inconsistent table columns near line {i}')
            make_table(document, rows)
            stats['tables'] += 1
            continue
        if text.startswith(('>', '<', '---', '~~~')):
            raise ValueError(f'Unsupported block syntax at line {i + 1}: {text}')
        ordered = re.match(r'^\d+\.\s+', text)
        unordered = re.match(r'^[-*]\s+', text)
        if ordered or unordered:
            prefix = ordered or unordered
            # Keep explicit numbers, so independent lists never accidentally continue.
            paragraph = document.add_paragraph(style='Normal')
            paragraph.paragraph_format.left_indent = Inches(.17)
            paragraph.paragraph_format.first_line_indent = Inches(-.17)
            inline(paragraph, text if ordered else '• ' + text[prefix.end():])
            i += 1
            continue
        pieces = [text]
        i += 1
        while i < len(lines) and lines[i].strip():
            if re.match(r'^(#|\||```|!\[|\d+\.\s|[-*]\s)', lines[i].strip()):
                break
            pieces.append(lines[i].strip())
            i += 1
        value = ' '.join(pieces)
        if introductory == 1:
            paragraph = document.add_paragraph(style='Subtitle')
            introductory = 2
        else:
            paragraph = document.add_paragraph()
        inline(paragraph, value)
        if re.match(r'^\*\*Table\s', value):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(4)
            for run in paragraph.runs:
                run.font.name, run.font.size = 'Calibri', Pt(10)
        elif re.match(r'^Figure\s+\d+\.', value):
            paragraph.style = 'Figure Caption'
        stats['paragraphs'] += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f'Created {output}: {stats}')


if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    convert(args.input.resolve(), args.output.resolve())
