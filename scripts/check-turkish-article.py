"""Audit the complete Turkish translation without altering either language.

Checks structural/numeric fidelity; semantic translation and visual pagination
still require human/visual review. Requires python-docx and lxml.
"""
from pathlib import Path
from zipfile import ZipFile
from collections import Counter
import hashlib, json, re
from lxml import etree
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
ARTICLE = ROOT / 'docs/article'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}


def tables(text):
    result, current = [], []
    for line in text.splitlines() + ['']:
        if line.startswith('|'):
            values = [v.strip() for v in line.strip('|').split('|')]
            if not all(re.fullmatch(':?-+:?', v) for v in values):
                current.append(values)
        elif current:
            result.append(current)
            current = []
    return result


def main():
    for name, digest in [
        ('morrow-article.md', '5215f423e991bc73edf7bcf329e089455c9fc8b19544e063fcee72c51cce8252'),
        ('morrow-article.docx', '28e3d3d7429d39ac9d312a14572ea9e82739ffe71805ca6f984b5e8dd67a4794')]:
        assert hashlib.sha256((ARTICLE / name).read_bytes()).hexdigest() == digest, name
    for item in json.loads((ARTICLE / 'artifact-manifest.json').read_text())['files']:
        content = (ROOT / item['path']).read_bytes()
        assert len(content) == item['bytes']
        assert hashlib.sha256(content).hexdigest() == item['sha256'], item['path']
    source = (ARTICLE / 'morrow-article.md').read_text(encoding='utf-8')
    translation = (ARTICLE / 'morrow-article.tr.md').read_text(encoding='utf-8')
    original, translated = source.splitlines(), translation.splitlines()
    assert len(original) == len(translated)
    # Every source paragraph, heading, table row, blank separator and comment
    # occupies its corresponding line. The actual prose is reviewed separately.
    def kind(line):
        if not line: return 'blank'
        if line.startswith('|'): return 'row'
        if line.startswith('#'): return line.split(' ')[0]
        if line.startswith('```'): return 'fence'
        if line.startswith('<!--'): return 'comment'
        if line.startswith('!['): return 'image'
        return 'text'
    number = re.compile(r'(?<![A-Za-z0-9])[+−-]?\d+(?:\.\d+)?(?:,\d{3})*(?![A-Za-z0-9])')
    for index, (a, b) in enumerate(zip(original, translated), 1):
        assert kind(a) == kind(b), f'Block changed at line {index}'
        # A cell may put a denominator first in Turkish ("0 of 12" -> "12
        # içinde 0"); compare per-cell number multisets, not prose word order.
        a_values = a.strip('|').split('|') if a.startswith('|') else [a]
        b_values = b.strip('|').split('|') if b.startswith('|') else [b]
        assert len(a_values) == len(b_values)
        for x, y in zip(a_values, b_values):
            assert Counter(number.findall(x)) == Counter(number.findall(y)), f'Numbers differ at line {index}: {x} / {y}'
    assert re.findall(r'```.*?```', source, re.S) == re.findall(r'```.*?```', translation, re.S)
    assert re.findall(r'https?://[^)\s`]+', source) == re.findall(r'https?://[^)\s`]+', translation)
    assert re.findall(r'(?<!`)`[^`\n]+`(?!`)', source) == re.findall(r'(?<!`)`[^`\n]+`(?!`)', translation)
    source_tables, tr_tables = tables(source), tables(translation)
    assert len(source_tables) == len(tr_tables) == 28
    document = Document(ARTICLE / 'morrow-article.tr.docx')
    assert len(document.tables) == 28
    for index, (expected, table, old) in enumerate(zip(tr_tables, document.tables, source_tables), 1):
        actual = [[c.text.replace('\n', '') for c in r.cells] for r in table.rows]
        assert expected == actual, f'Turkish Word table {index} differs from Markdown'
        assert [len(r) for r in old] == [len(r) for r in expected]
        assert table.rows[0]._tr.trPr.find('{' + NS['w'] + '}tblHeader') is not None
        for a_row, b_row in zip(old, expected):
            for a, b in zip(a_row, b_row):
                if re.fullmatch(r'[+−-]?\d+(?:\.\d+)?|[0-9a-f]{64}', a):
                    assert a == b, f'Exact numeric/hash cell changed in table {index}'
                if a in ['Akım', 'Çiy', 'Köz', 'Sis', 'Yankı', 'Filiz', 'Köz, Filiz', 'Yankı, Filiz', 'A', 'D', 'XA', 'XD']:
                    assert a == b
    assert sum(len(t) == 101 for t in tr_tables) == 4
    with ZipFile(ARTICLE / 'morrow-article.docx') as old, ZipFile(ARTICLE / 'morrow-article.tr.docx') as new:
        assert old.namelist() == new.namelist(), 'Package part added/lost'
        changed = [name for name in old.namelist() if old.read(name) != new.read(name)]
        allowed = ['word/document.xml', 'docProps/core.xml'] + [f'word/media/image{i}.png' for i in range(1, 4)]
        assert set(changed) == set(allowed), changed
        a = etree.fromstring(old.read('word/document.xml'))
        b = etree.fromstring(new.read('word/document.xml'))
        math_a = a.xpath('//m:oMath', namespaces=NS)
        math_b = b.xpath('//m:oMath', namespaces=NS)
        assert len(math_a) == len(math_b) == 10
        assert [etree.tostring(x) for x in math_a] == [etree.tostring(x) for x in math_b]
        assert len(b.xpath('//w:drawing', namespaces=NS)) == 3
        for left, right in zip(a.xpath('//w:tblGrid | //w:tcPr | //w:sectPr', namespaces=NS), b.xpath('//w:tblGrid | //w:tcPr | //w:sectPr', namespaces=NS)):
            assert etree.tostring(left) == etree.tostring(right), 'Template geometry changed'
        # Cross-check every top-level prose paragraph, including hyperlink text,
        # against the converter's independently assembled translation content.
        converter_output = ROOT / '.local/translation-qa/translated-content.docx'
        if converter_output.is_file():
            with ZipFile(converter_output) as content:
                expected = etree.fromstring(content.read('word/document.xml'))
                text = lambda root: [''.join(p.xpath('.//w:t/text()', namespaces=NS)) for p in root.xpath('./w:body/w:p', namespaces=NS)]
                assert text(b) == text(expected)
    images = re.findall(r'!\[[^]]*\]\(([^)]+)\)', translation)
    assert len(images) == 3
    for index, path in enumerate(images, 1):
        assert '.tr.png' in path
        with ZipFile(ARTICLE / 'morrow-article.tr.docx') as package:
            assert package.read(f'word/media/image{index}.png') == (ARTICLE / path).read_bytes()
        # Check plotted marks independently of translated labels/font metrics.
        original_svg = etree.parse(str(ARTICLE / path.replace('.tr.png', '.svg')))
        translated_svg = etree.parse(str(ARTICLE / path.replace('.png', '.svg')))
        marks = lambda svg: [(etree.QName(x).localname, dict(x.attrib)) for x in svg.iter() if etree.QName(x).localname == 'circle']
        assert marks(original_svg) == marks(translated_svg), 'Plotted data points changed'
    manifest = ARTICLE / 'turkish-translation-manifest.json'
    if manifest.exists():
        for item in json.loads(manifest.read_text())['files']:
            data = (ROOT / item['path']).read_bytes()
            assert len(data) == item['bytes'] and hashlib.sha256(data).hexdigest() == item['sha256'], item['path']
    print(f'Turkish translation audit passed: {len(original)} aligned lines, 28 tables, four complete 100-seed ledgers, 3 translated figures, 10 unchanged native equations; English and scientific artifacts unchanged.')


if __name__ == '__main__':
    main()
