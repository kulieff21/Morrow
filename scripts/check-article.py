"""Audit the article's table values and DOCX structure against preserved inputs.

Requires python-docx. Does not replace visual page inspection or model replay.
"""
from pathlib import Path
import hashlib
import json
import re
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
ARTICLE = ROOT / 'docs/article/morrow-article.md'
DOCX = ARTICLE.with_suffix('.docx')


def rows(markdown):
    tables, current = [], []
    for line in markdown.splitlines() + ['']:
        if line.startswith('|'):
            row = [value.strip() for value in line.strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?', value) for value in row):
                current.append(row)
        elif current:
            tables.append(current)
            current = []
    return tables


def main():
    text = ARTICLE.read_text(encoding='utf-8')
    tables = rows(text)
    document = Document(DOCX)
    assert len(tables) == len(document.tables) == 21
    for index, (expected, table) in enumerate(zip(tables, document.tables), 1):
        actual = [[cell.text.replace('\n', '') for cell in row.cells] for row in table.rows]
        assert actual == expected, f'DOCX table {index} differs from Markdown'
        assert table.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
    paired = json.loads((ROOT / 'docs/paired-analysis.json').read_text(encoding='utf-8'))
    thrift = json.loads((ROOT / 'docs/thrift-analysis.json').read_text(encoding='utf-8'))
    for source in [paired['source'], thrift['source']]:
        assert hashlib.sha256((ROOT / source['path']).read_bytes()).hexdigest() == source['sha256']
    # Check hand-written summary tables directly against machine-readable analysis.
    for row in paired['lineages']:
        expected = [row['name'], f"{row['meanControlCount']:.2f}", f"{row['meanDroughtCount']:.2f}",
                    f"{100 * row['meanControlShare']:.2f}", f"{100 * row['meanDroughtShare']:.2f}",
                    f"{row['meanShareDeltaPercentagePoints']:+.2f}".replace('-', '−')]
        assert expected in tables[4], f'Lineage mean differs for {row["name"]}'
    for label, key in [('A', 'naturalAbundance'), ('D', 'naturalDrought'),
                       ('XA', 'normalizedAbundance'), ('XD', 'normalizedDrought')]:
        arm = thrift['arms'][key]
        expected = [label, f"{arm['meanPopulation']:.2f}", f"{arm['meanCount']:.2f}",
                    f"{arm['meanSharePercent']:.2f}", f"{arm['targetExtinctions']} of 12"]
        assert expected in tables[5]
    for outcome, index in [('share', 6), ('count', 7)]:
        for key, label in [('abundance', 'XA minus A'), ('drought', 'XD minus D'), ('interaction', 'Interaction')]:
            value = thrift['effects'][outcome][key]
            fmt = lambda x: f'{x:+.2f}'.replace('-', '−') if x > 0 else f'{x:.2f}'.replace('-', '−')
            expected = [label, fmt(value['mean']), fmt(value['median'])] + [fmt(x) for x in value['range']] + [str(value[k]) for k in ['negative', 'zero', 'positive']]
            assert expected in tables[index], (outcome, key, expected)
    with ZipFile(DOCX) as archive:
        xml = archive.read('word/document.xml').decode('utf-8')
        assert xml.count('<m:oMath>') == 10
        assert xml.count('<m:sSup>') == 2
        assert '<m:sSub>' in xml
        assert xml.count('<w:drawing>') == 2
        assert '<w:pBdr>' not in xml
        assert 'BEGIN GENERATED' not in xml
        assert all(name in xml for name in ['Çiy', 'Akım', 'Yankı'])
        rels = archive.read('word/_rels/document.xml.rels').decode('utf-8')
        assert 'https://doi.org/10.18564/jasss.4259' in rels
    for name in ['src/engine.js', 'docs/initial-batch.json', 'docs/thrift-protocol.md', 'docs/thrift-raw.json']:
        digest = hashlib.sha256((ROOT / name).read_bytes()).hexdigest()
        assert digest in text, f'Frozen source hash not present: {name}'
    print('Article audit passed: 21 tables, summary values, 10 native equations, 2 figures, links and frozen hashes.')


if __name__ == '__main__':
    main()
